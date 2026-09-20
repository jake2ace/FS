"""Turn an attachment (txt / pdf / docx / xlsx) into plain text lines and detect
what kind of shipping document it is.

Design rule: never invent content. If a file cannot be read reliably, the
document is marked unreadable and the case goes to human review.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import Optional

DOC_SI = "SI"
DOC_BL = "BL"
DOC_CI = "COMMERCIAL_INVOICE"
DOC_PL = "PACKING_LIST"
DOC_COO = "CERTIFICATE_OF_ORIGIN"
DOC_UNKNOWN = "UNKNOWN"

DOC_TYPE_LABELS = {
    DOC_SI: "Shipping Instruction",
    DOC_BL: "Draft Bill of Lading",
    DOC_CI: "Commercial Invoice",
    DOC_PL: "Packing List",
    DOC_COO: "Certificate of Origin",
    DOC_UNKNOWN: "Unknown document",
}


@dataclass
class ParsedDoc:
    path: str
    filename: str
    fmt: str
    size: int
    readable: bool = True
    error: Optional[str] = None
    lines: list[str] = field(default_factory=list)
    text: str = ""
    detected_type: str = DOC_UNKNOWN
    role_hint: Optional[str] = None


def role_hint_from_name(filename: str) -> Optional[str]:
    stem = filename.rsplit(".", 1)[0].upper()
    if stem.endswith("_SI") or stem.endswith("-SI") or stem.endswith(" SI"):
        return DOC_SI
    if stem.endswith("_BL") or stem.endswith("-BL") or stem.endswith(" BL"):
        return DOC_BL
    return None


def _fmt(filename: str) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    return ext if ext in {"txt", "pdf", "docx", "xlsx"} else (ext or "other")


def _cell_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


# ---------------------------------------------------------------------------
# readers
# ---------------------------------------------------------------------------

def _read_txt(data: bytes) -> list[str]:
    text = data.decode("utf-8", errors="replace")
    return text.splitlines()


def _join_chars(chars: list[dict]) -> str:
    """Join pdfplumber chars left-to-right, inserting a space at visible gaps."""
    chars = sorted(chars, key=lambda c: c["x0"])
    out: list[str] = []
    prev = None
    for c in chars:
        if prev is not None and c["text"] != " " and (c["x0"] - prev["x1"]) > 1.0 and out and not out[-1].endswith(" "):
            out.append(" ")
        out.append(c["text"])
        prev = c
    return "".join(out)


def _pdf_lines_from_chars(page) -> list[str]:
    """Layout-aware line builder.

    Form-style PDFs print the label in bold and the value in a regular font on
    the same baseline; long labels can physically overlap the value column,
    which makes plain text extraction interleave the two strings.  Splitting a
    line by font weight recovers 'Label: value' cleanly.
    """
    chars = [c for c in page.chars if c.get("text") is not None]
    if not chars:
        return []
    chars.sort(key=lambda c: (c["top"], c["x0"]))
    groups: list[list[dict]] = []
    for c in chars:
        if groups and abs(c["top"] - groups[-1][0]["top"]) <= 2.5:
            groups[-1].append(c)
        else:
            groups.append([c])
    lines: list[str] = []
    for g in groups:
        bold = [c for c in g if "bold" in str(c.get("fontname", "")).lower()]
        regular = [c for c in g if "bold" not in str(c.get("fontname", "")).lower()]
        if bold and regular:
            label = _join_chars(bold).strip().rstrip(":：").strip()
            value = _join_chars(regular).strip()
            lines.append(f"{label}: {value}" if label else value)
        else:
            lines.append(_join_chars(g).strip())
    return lines


def _read_pdf(data: bytes) -> tuple[list[str], Optional[str]]:
    """Returns (lines, error). Image-only PDFs come back with an error."""
    import pdfplumber  # local import keeps startup fast

    lines: list[str] = []
    n_images = 0
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        if not pdf.pages:
            return [], "PDF has no pages"
        for page in pdf.pages:
            n_images += len(page.images or [])
            page_lines: list[str] = []
            try:
                page_lines = _pdf_lines_from_chars(page)
            except Exception:
                page_lines = []
            if not any(l.strip() for l in page_lines):
                txt = page.extract_text() or ""
                page_lines = txt.splitlines()
            lines.extend(page_lines)
    if not any(l.strip() for l in lines):
        if n_images:
            return [], "image-only PDF (no text layer); OCR is not configured"
        return [], "PDF contains no extractable text"
    return lines, None


def _read_docx(data: bytes) -> list[str]:
    import docx  # python-docx

    document = docx.Document(io.BytesIO(data))
    lines: list[str] = []
    for para in document.paragraphs:
        if para.text.strip():
            lines.append(para.text.strip())
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            # de-duplicate merged cells that python-docx repeats
            dedup: list[str] = []
            for c in cells:
                if not dedup or dedup[-1] != c:
                    dedup.append(c)
            cells = [c for c in dedup if c]
            if not cells:
                continue
            if len(cells) == 1:
                lines.append(cells[0])
            else:
                label = cells[0].replace("\n", " ")
                value = " | ".join(x.strip() for x in cells[1].split("\n") if x.strip())
                rest = " | ".join(cells[2:])
                lines.append(f"{label}: {value}" + (f" | {rest}" if rest else ""))
    return lines


def _read_xlsx(data: bytes) -> list[str]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    lines: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [_cell_text(c) for c in row]
            cells = [c for c in cells if c]
            if not cells:
                continue
            if len(cells) == 1:
                lines.append(cells[0])
            else:
                lines.append(f"{cells[0]}: {cells[1]}" + (" | " + " | ".join(cells[2:]) if len(cells) > 2 else ""))
    return lines


# ---------------------------------------------------------------------------
# document type detection
# ---------------------------------------------------------------------------

def detect_doc_type(lines: list[str], filename: str = "") -> str:
    head = " ".join(l.strip() for l in lines[:6] if l.strip()).upper()
    head = re.sub(r"\s+", " ", head)[:400]
    if "COMMERCIAL INVOICE" in head:
        return DOC_CI
    if "PACKING LIST" in head:
        return DOC_PL
    if "CERTIFICATE OF ORIGIN" in head:
        return DOC_COO
    if "INSTRUCTION" in head or "S.I." in head or re.search(r"\bSI\b", head):
        return DOC_SI
    if "BILL OF LADING" in head or re.search(r"\bB/L\b|\bBL\b", head):
        return DOC_BL
    # fall back to the whole text if the title line is unusual
    body = " ".join(lines).upper()
    if "SHIPPING INSTRUCTION" in body:
        return DOC_SI
    if "BILL OF LADING" in body:
        return DOC_BL
    return DOC_UNKNOWN


# ---------------------------------------------------------------------------
# public entry point
# ---------------------------------------------------------------------------

def parse_attachment(path: str, data: bytes) -> ParsedDoc:
    filename = path.rsplit("/", 1)[-1]
    doc = ParsedDoc(path=path, filename=filename, fmt=_fmt(filename), size=len(data),
                    role_hint=role_hint_from_name(filename))
    if len(data) == 0:
        doc.readable = False
        doc.error = "empty file (0 bytes)"
        return doc
    try:
        if doc.fmt == "txt":
            doc.lines = _read_txt(data)
        elif doc.fmt == "pdf":
            doc.lines, err = _read_pdf(data)
            if err:
                doc.readable = False
                doc.error = err
        elif doc.fmt == "docx":
            doc.lines = _read_docx(data)
        elif doc.fmt == "xlsx":
            doc.lines = _read_xlsx(data)
        else:
            # unknown extension: try as text if it decodes cleanly
            try:
                txt = data.decode("utf-8")
                doc.lines = txt.splitlines()
            except UnicodeDecodeError:
                doc.readable = False
                doc.error = f"unsupported attachment format '{doc.fmt}'"
    except Exception as exc:  # corrupt / truncated / not really that format
        doc.readable = False
        doc.error = f"could not open {doc.fmt.upper()} file: {type(exc).__name__}: {str(exc)[:120]}"

    if doc.readable and not any(l.strip() for l in doc.lines):
        doc.readable = False
        doc.error = doc.error or "file contains no readable text"

    doc.text = "\n".join(doc.lines)
    if doc.readable:
        doc.detected_type = detect_doc_type(doc.lines, filename)
    return doc
