"""Regression coverage for the confirmed workflow; no live AI calls."""
import base64
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from app.ai import AIClient
from tests.fakes import FixtureAI
from app.pipeline import Analyser
from app.schemas import FIELDS, FieldValue
from app.extract import valid_value
from app.revisions import build_revision, revision_path
from app.store import Store
from app.submission import to_entry

TEXT = '''Shipper: ACME EXPORTS
Consignee: BUYER COMPANY
Notify Party: BUYER COMPANY
Port of Loading: SINGAPORE
Port of Discharge: PORT KLANG
Container Count: 3
Gross Weight: 22000 KG
'''
EMAIL = dict(email_id='email_test', subject='Draft BL for checking', body='Please compare the SI and draft BL. Attached are the SI and draft BL.',
             **{'from': 'ops@example.com'}, attachments=['attachments/test_SI.txt', 'attachments/test_BL.txt'])

class MemoryInbox:
    def __init__(self, docs): self.docs = docs
    def read_bytes(self, path):
        if path not in self.docs: raise FileNotFoundError(path)
        value = self.docs[path]
        return value.encode() if isinstance(value, str) else value
    def get(self, eid): return dict(EMAIL) if eid == EMAIL['email_id'] else None
    def emails(self): return [dict(EMAIL)]
    def __len__(self): return 1
    def attachment_size(self, path): return len(self.read_bytes(path))


def inputs(text=TEXT):
    return {EMAIL['attachments'][0]: 'SHIPPING INSTRUCTION\n' + text,
            EMAIL['attachments'][1]: 'DRAFT BILL OF LADING\n' + text}


def off(): return FixtureAI()

async def test_no_attachment_even_request_draft_is_review():
    r = await Analyser(MemoryInbox({}), off()).analyse(dict(EMAIL, attachments=[], body='Please send the draft BL.'))
    assert (r.status, r.review_reason, r.automation) == ('NEEDS_REVIEW', 'missing_attachment', 'review_required')


async def test_missing_physical_file_is_missing_attachment():
    d = inputs(); del d[EMAIL['attachments'][1]]
    r = await Analyser(MemoryInbox(d), off()).analyse(EMAIL)
    assert r.review_reason == 'missing_attachment'


async def test_unknown_titles_cannot_be_trusted_from_filenames():
    r = await Analyser(MemoryInbox({p: TEXT for p in EMAIL['attachments']}), off()).analyse(EMAIL)
    assert r.status == 'NEEDS_REVIEW' and r.review_reason == 'wrong_doc_type'


@pytest.mark.parametrize('field,value', [('container_count','3.5'),('container_count','3 or 4'),('container_count','0'),('gross_weight_kg','-30 KG'),('gross_weight_kg','20 bananas'),('gross_weight_kg','NaN')])
def test_ambiguous_or_invalid_numbers(field, value):
    assert not valid_value(field, value)


async def test_reanalysis_reopens_instead_of_inheriting_old_resolution():
    store = Store()
    analyser = Analyser(MemoryInbox(inputs()), off())
    r = await analyser.analyse(EMAIL); store.put(r)
    store.record_decision(r.email_id, 'confirm', 'Checked')
    data = inputs(); data[EMAIL['attachments'][1]] = data[EMAIL['attachments'][1]].replace('Count: 3','Count: 4')
    r2 = await Analyser(MemoryInbox(data), off()).analyse(EMAIL); store.put(r2)
    assert r2.status == 'MISMATCH' and not r2.resolved and r2.decision is None
    assert r2.history[-1]['previous']['resolved']


@pytest.fixture
def client(tmp_path, monkeypatch):
    from app import main
    data = inputs(); data[EMAIL['attachments'][1]] = data[EMAIL['attachments'][1]].replace('Count: 3', 'Count: 4')
    ib = MemoryInbox(data)
    monkeypatch.setattr(main, 'inbox', ib)
    monkeypatch.setattr(main, 'store', Store(tmp_path))
    monkeypatch.setattr(main, 'ai', off())
    monkeypatch.setattr(main, 'analyser', Analyser(ib, main.ai))
    monkeypatch.setattr(main.config, 'CACHE_DIR', tmp_path)
    monkeypatch.setattr(main, '_case_locks', {})
    c = TestClient(main.app)
    r = c.post('/api/analyse/email_test?explain=false')
    assert r.status_code == 200 and r.json()['status'] == 'MISMATCH'
    return c, main, tmp_path


@pytest.mark.parametrize('action,retained', [('approve_revision', True), ('reject_revision', False)])
def test_native_copy_adoption_or_rejection(client, action, retained):
    c, main, tmp = client
    before = main.inbox.read_bytes(EMAIL['attachments'][1])
    original = c.get('/api/results/email_test').json()
    response = c.post('/api/cases/email_test/revision')
    assert response.status_code == 200, response.text
    r = response.json(); report = r['working_report']
    assert r['processing_status'] == 'PENDING_HUMAN_APPROVAL' and not r['resolved']
    path = revision_path(tmp, 'email_test', report)
    assert b'Count: 3' in path.read_bytes() and b'Count: 4' in before
    assert c.post('/api/cases/email_test/decision', json={'action':'approve_revision', 'revision_id':'stale'}).status_code == 409
    assert c.post('/api/analyse/email_test?explain=false').status_code == 409
    response = c.post('/api/cases/email_test/decision', json={'action':action, 'revision_id':report['revision_id']})
    assert response.status_code == 200, response.text
    assert path.exists() is retained
    assert response.json()['resolved'] is retained
    assert main.inbox.read_bytes(EMAIL['attachments'][1]) == before
    assert response.json()['status'] == original['status'] == 'MISMATCH'
    assert c.get('/api/submission').json()['email_test']['defect_fields'] == ['container_count']
    assert c.get(f"/api/cases/email_test/revision/{report['revision_id']}/file").status_code == (200 if retained else 404)
    restored = Store(tmp); restored.load()
    assert restored.get('email_test').resolved is retained
    assert restored.get('email_test').history


def test_mismatch_cannot_be_closed_without_review(client):
    c, _, _ = client
    assert c.post('/api/cases/email_test/decision', json={'action':'confirm'}).status_code == 409


def test_human_readings_and_upload_update_only_working_report(client):
    c, _, _ = client
    old = c.get('/api/results/email_test').json()
    fields = {row['field']: {'si_value':row['si_value'],'bl_value':row['si_value']} for row in old['fields']}
    assert c.post('/api/cases/email_test/readings', json={'fields':fields}).status_code == 422
    r = c.post('/api/cases/email_test/readings', json={'fields':fields,'note':'Read the original sources; corrected OCR reading.'}).json()
    assert r['status'] == 'MISMATCH' and r['working_report']['status'] == 'OK'
    c.post('/api/cases/email_test/decision', json={'action':'reject_revision','revision_id':r['working_report']['revision_id']})
    uploads = [{'name':Path(p).name,'base64':base64.b64encode(text.encode()).decode()} for p,text in inputs().items()]
    response = c.post('/api/cases/email_test/attachments', json={'files':uploads,'note':'Carrier supplied corrected documents.'})
    assert response.status_code == 200, response.text
    r = response.json()
    assert r['status'] == 'MISMATCH' and r['working_report']['status'] == 'OK'
    assert len(r['history']) >= 4


def test_retry_only_failed_items_and_no_double_count(client):
    c, main, _ = client
    run = main.store.new_run(1, True); run.status = 'completed'; run.done = 1
    run.failed = [{'email_id':'email_test','error':'network'}]
    assert c.post(f'/api/runs/{run.run_id}/retry/email_test').status_code == 200
    assert c.post(f'/api/runs/{run.run_id}/retry/email_test').status_code == 409
    assert run.done == 1 and run.mismatch == 1
    assert c.post(f'/api/runs/{run.run_id}/cancel').status_code == 409


async def test_native_docx_xlsx_pdf_copies_reparse_and_keep_format(tmp_path):
    import io
    import pymupdf
    from docx import Document
    from openpyxl import Workbook
    for fmt in ['docx','xlsx','pdf']:
        body = 'DRAFT BILL OF LADING\n' + TEXT.replace('Count: 3','Count: 4')
        out = io.BytesIO()
        if fmt == 'docx':
            doc = Document()
            for line in body.splitlines(): doc.add_paragraph(line)
            doc.save(out); data=out.getvalue()
        elif fmt == 'xlsx':
            wb=Workbook();ws=wb.active;ws.append(['DRAFT BILL OF LADING'])
            for line in body.splitlines()[1:]: ws.append(line.split(': ',1))
            wb.save(out);data=out.getvalue()
        else:
            doc=pymupdf.open();page=doc.new_page()
            page.insert_text((40,40),'DRAFT BILL OF LADING',fontsize=12)
            for i,line in enumerate(body.splitlines()[1:]):
                label,value=line.split(': ',1)
                page.insert_text((40,80+i*30),label+':',fontsize=11)
                page.insert_text((220,80+i*30),value,fontsize=11)
            data=doc.tobytes()
        path='attachments/test_BL.'+fmt
        ib=MemoryInbox({EMAIL['attachments'][0]:'SHIPPING INSTRUCTION\n'+TEXT,path:data})
        r=await Analyser(ib,off()).analyse(dict(EMAIL,attachments=[EMAIL['attachments'][0],path]))
        assert r.status == 'MISMATCH', (fmt,r.model_dump())
        report=await build_revision(r,ib,tmp_path,Analyser(ib,off()))
        copy=revision_path(tmp_path,r.email_id,report)
        assert copy.suffix == '.'+fmt and copy.read_bytes() != data
        assert report['status']=='OK' and all(row['match'] for row in report['fields'])
        assert ib.read_bytes(path)==data


def test_small_weight_and_similar_company_are_real_differences():
    from app.compare import parties_equal, weight_equal
    assert weight_equal('22000 KG','22000.5 KG')[0] is False
    assert parties_equal('ACME EXPORTS 123 LIMITED','ACME EXPORTS 124 LIMITED')[0] is False
    assert parties_equal('ACME LIMITED','ACME LTD')[0] is True
    assert parties_equal('ACME LLC','ACME LTD')[0] is False


def test_real_local_ocr_recovers_scanned_pdf():
    import shutil
    import pymupdf
    from app.parsers import parse_attachment
    from app.recovery import recover_document
    from app.extract import extract_fields, missing_fields
    if not shutil.which('tesseract') or not shutil.which('pdftoppm'):
        pytest.skip('Local OCR tools not installed')
    source=pymupdf.open(); page=source.new_page()
    for i,line in enumerate(('DRAFT BILL OF LADING\n'+TEXT).splitlines()):
        page.insert_text((40,50+i*35),line,fontsize=15)
    image=page.get_pixmap(matrix=pymupdf.Matrix(2,2)).tobytes('png')
    scan=pymupdf.open();page=scan.new_page();page.insert_image(page.rect,stream=image)
    data=scan.tobytes()
    parsed=parse_attachment('scan_BL.pdf',data)
    assert not parsed.readable
    recovered=recover_document(parsed,data)
    assert recovered.readable and recovered.recovery=='ocr'
    assert recovered.detected_type=='BL'
    assert not missing_fields(extract_fields(recovered))


def test_old_cache_archived_and_interrupted_run_visible(tmp_path):
    import json
    from app.schemas import CaseResult
    old=CaseResult(email_id='old',subject='old',sender='test',category='GENERAL',category_confidence=.9,status='OK',analysed_at='2026-09-20')
    store=Store(tmp_path);store.results['old']=old;store.new_run(20,False);store.flush(force=True)
    restored=Store(tmp_path);restored.load()
    assert restored.get('old') is None and len(restored.legacy_results)==1
    assert restored.runs_list()[0].status=='failed'
    assert 'restarted' in restored.runs_list()[0].failed[0]['error']


def test_changed_copy_cannot_be_adopted(client):
    c, main, tmp = client
    report=c.post('/api/cases/email_test/revision').json()['working_report']
    path=revision_path(tmp,'email_test',report)
    path.write_bytes(path.read_bytes()+b'\nchanged after review')
    r=c.post('/api/cases/email_test/decision',json={'action':'approve_revision','revision_id':report['revision_id']})
    assert r.status_code==409 and 'changed' in r.json()['detail']
    assert not main.store.get('email_test').resolved


def test_export_requires_all_emails_analysed(client):
    c,main,_=client
    main.store.results.clear()
    r=c.get('/api/submission')
    assert r.status_code==409 and 'remaining 1' in r.json()['detail']


def test_adopted_copy_ends_ai_processing_and_stays_downloadable(client):
    c,main,tmp=client
    report=c.post('/api/cases/email_test/revision').json()['working_report']
    assert c.post('/api/cases/email_test/decision',json={'action':'approve_revision','revision_id':report['revision_id']}).status_code==200
    assert c.post('/api/analyse/email_test?explain=false').status_code==409
    assert main.store.get('email_test').resolved
    assert revision_path(tmp,'email_test',report).exists()
    assert c.get(f"/api/cases/email_test/revision/{report['revision_id']}/file").status_code==200
